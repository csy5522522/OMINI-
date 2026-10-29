"""Import row-based Excel/CSV timetables into the existing schedule schema."""
import base64
import csv
import io

from school_timetable import parse_schedule

HEADERS = ('课程名称', '星期', '节次', '周次', '教师', '教室', '校区', '学分')


def import_table(content, filename, year, term):
    try:
        if filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            import zipfile
            raw = base64.b64decode(content, validate=True)
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 48 * 1024 * 1024:
                    raise ValueError('表格解压后超过 48 MB，请精简后导入。')
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                rows = list(_bounded(book.active.iter_rows(values_only=True)))
            finally:
                book.close()
        elif filename.lower().endswith(('.csv', '.tsv')):
            rows = list(_bounded(csv.reader(io.StringIO(content.lstrip('\ufeff')), delimiter='\t' if filename.lower().endswith('.tsv') else ',')))
        else:
            raise ValueError('请选择 XLSX、CSV 或 TSV 表格。')
        if not rows:
            raise ValueError('表格为空。')
        headers = [str(v or '').strip() for v in rows[0]]
        missing = [name for name in HEADERS[:4] if name not in headers]
        if missing:
            raise ValueError('表格缺少列：' + '、'.join(missing) + '。请下载模板填写。')
        courses = []
        for index, row in enumerate(rows[1:], 2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            values = dict(zip(headers, (str(v).strip() if v is not None else '' for v in row)))
            day = values.get('星期', '').replace('星期', '').replace('周', '')
            day = str({'一':1, '二':2, '三':3, '四':4, '五':5, '六':6, '日':7, '天':7}.get(day, day))
            course = dict(zip(('kcmc', 'xqj', 'jcs', 'zcd', 'xm', 'cdmc', 'xqmc', 'xf'), (values.get(h, '') for h in HEADERS)))
            course['xqj'] = day
            if not course['kcmc']:
                raise ValueError(f'第 {index} 行缺少课程名称。')
            try:
                parse_schedule({'kbList': [course]}, str(year), str(term))
            except ValueError as error:
                raise ValueError(f'第 {index} 行：{error}') from error
            courses.append(course)
        if not courses:
            raise ValueError('表格中没有课程。')
        return parse_schedule({'kbList': courses}, str(year), str(term))
    except ImportError as error:
        raise ValueError('Excel 组件尚未安装，请安装 requirements-web.txt 中的依赖或使用 CSV 表格。') from error
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('表格读取失败，请检查文件格式后重试。') from error


def _bounded(rows):
    for index, row in enumerate(rows):
        if index > 5000 or len(row) > 64:
            raise ValueError('表格最多支持 5000 条课程和 64 列。')
        yield row
