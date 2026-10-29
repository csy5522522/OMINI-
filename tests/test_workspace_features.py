"""Local integration checks for streaming, history, previews and public caches."""
import base64
import asyncio
from io import BytesIO
import json
from pathlib import Path
import secrets
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

import requests
import httpx
from PIL import Image
from openpyxl import Workbook
from account_client import AccountClient
from accounts import AccountStore
from account_data import AccountDataStore
from chat_history import ChatHistory, ConversationConflict, HistoryQuota
from timetable_ai import TimetableAIService, normalize_preview
from trends import TrendsService, _items, _baidu_rows, SOURCES
from chat_runtime import stream_model, ChatBusy
from web_chat import ChatService
import tests.test_web_chat as fixtures


def file(name, raw):return {'name':name,'data':base64.b64encode(raw).decode()}


class StreamHTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):fixtures.WebChatTests.setUpClass.__func__(cls)
    @classmethod
    def tearDownClass(cls):fixtures.WebChatTests.tearDownClass.__func__(cls)

    def test_stream_history_idempotency_cancel_and_isolation(self):
        admin=AccountClient(self.url);admin.request('login',{'email':'admin','password':'admin123'})
        user=AccountClient(self.url);created=user.request('register',{'email':'stream@example.test','password':'FixturePass-123'})['user']
        other=AccountClient(self.url);other.request('register',{'email':'other-stream@example.test','password':'FixturePass-123'})
        s=requests.Session();s.trust_env=False;s.cookies.set('omini_session',user.token);s.headers['X-Omini-Client']='web';self.addCleanup(s.close)
        identities=[]
        for protocol in ('openai','responses','anthropic'):
            config={'label':protocol,'protocol':protocol,'base_url':self.base,'api_key':'fixture-secret-key','models':[{'name':'fixture-model','alias':'fixture'}]}
            identity=admin.api('/api/admin/chat/providers',config)['id'];identities.append(identity)
            body={'model':identity+':fixture','messages':[{'role':'user','content':'中文分片加密历史'}],'request_id':secrets.token_hex(12),'version':0}
            start=time.monotonic();arrivals=[];events=[]
            def consume(payload):
                with s.post(self.url+'/api/chat/stream',json=payload,stream=True,timeout=15) as response:
                    response.raise_for_status();self.assertIn('text/event-stream',response.headers['Content-Type'])
                    self.assertEqual(response.headers['X-Accel-Buffering'],'no')
                    for line in response.iter_lines(chunk_size=1):
                        if line.startswith(b'data: '):
                            data=json.loads(line[6:]);events.append(data)
                            if 'text' in data:arrivals.append(time.monotonic()-start)
            before=len(fixtures.Upstream.calls);consume(body)
            self.assertLess(arrivals[0],arrivals[-1]);self.assertEqual(len(fixtures.Upstream.calls),before+1)
            conversation=events[-1]['conversation'];self.assertEqual(conversation['status'],'complete')
            self.assertEqual(conversation['messages'][-1]['content'],'你好，测试回复')
            cid=conversation['id'];events=[];consume(body);self.assertEqual(len(fixtures.Upstream.calls),before+1)
            with self.assertRaises(ValueError):other.api('/api/chat/conversations/'+cid)
            with self.assertRaises(ValueError):user.api('/api/chat/conversations/'+cid,{'title':'old','version':0})
            detail=user.api('/api/chat/conversations/'+cid,{'title':'已改名','version':conversation['version']})['conversation']
            body.update(conversation_id=cid,version=detail['version'],request_id=secrets.token_hex(12),messages=[{'role':'user','content':'SLOW'}])
            run_id=None
            with s.post(self.url+'/api/chat/stream',json=body,stream=True,timeout=15) as response:
                for line in response.iter_lines(chunk_size=1):
                    if line.startswith(b'data: '):
                        data=json.loads(line[6:]);run_id=data.get('run_id') or run_id
                        if data.get('text'):
                            user.api('/api/chat/runs/'+run_id+'/cancel',{})
                            break
            saved=user.api('/api/chat/conversations/'+cid)['conversation']
            self.assertEqual(saved['status'],'cancelled');self.assertEqual(saved['messages'][-1]['content'],'部分回复')
            self.assertEqual(user.api('/api/chat/conversations')['conversations'][0]['id'],cid)
            user.api('/api/chat/conversations/'+cid,{'action':'delete','version':saved['version']})
        # History deletion is included in user deletion.
        body={'model':identities[0]+':fixture','messages':[{'role':'user','content':'delete-account-history'}],'request_id':secrets.token_hex(12),'version':0};events=[];consume(body)
        admin.api('/api/admin/user',{'id':created['id'],'action':'delete'})
        with AccountStore(self.root/'server'/'accounts.sqlite3').connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) AS n FROM chat_conversations WHERE user_id=?',(created['id'],)).fetchone()['n'],0)
        for identity in identities:admin.api('/api/admin/chat/providers',{'action':'delete','id':identity})
        self.assertNotIn('中文分片加密历史'.encode(),(self.root/'server'/'accounts.sqlite3').read_bytes())

    def test_ai_http_preview_never_writes_and_permissions(self):
        admin=AccountClient(self.url);admin.request('login',{'email':'admin','password':'admin123'})
        member=AccountClient(self.url);member.request('register',{'email':'preview-http@example.test','password':'FixturePass-123'})
        config={'label':'识别测试','protocol':'openai','base_url':self.base,'api_key':'fixture-key','models':[{'name':'fixture-model','alias':'fixture'}],'default':True}
        identity=admin.api('/api/admin/timetable-ai/providers',config)['id']
        self.assertNotIn('api_key',json.dumps(member.api('/api/timetable/ai/routes')))
        with self.assertRaises(ValueError):member.api('/api/admin/timetable-ai/providers')
        stream=BytesIO();Image.new('RGB',(8,8),'white').save(stream,'PNG');files=[file('课表.png',stream.getvalue())]
        before=member.api('/api/timetable')
        preview=member.api('/api/timetable/ai/recognize',{'year':'2026','term':'3','model':identity,'files':files})
        self.assertFalse(preview['ready']);self.assertEqual(member.api('/api/timetable'),before)
        preview['courses'][0]['weeks']='1-16'
        preview=member.api('/api/timetable/ai/validate',preview);self.assertTrue(preview['ready'])
        saved=member.api('/api/timetable/import',{'content':json.dumps(preview['schedule']),'base_revision':before['revision']})
        self.assertEqual(saved['schedule']['courses'][0]['name'],'高等数学')
        changed=json.loads(json.dumps(preview['schedule']));changed['courses'][0]['name']='other device'
        latest=member.api('/api/timetable/import',{'content':json.dumps(changed),'base_revision':saved['revision']})
        preview['schedule']['courses'][0]['name']='local preview edit'
        conflict=member.api('/api/timetable/import',{'content':json.dumps(preview['schedule']),'base_revision':saved['revision']})
        self.assertTrue(conflict['conflicts']);self.assertEqual(member.api('/api/timetable')['schedule']['courses'][0]['name'],latest['schedule']['courses'][0]['name'])
        admin.api('/api/admin/timetable-ai/providers',dict(admin.api('/api/admin/timetable-ai/providers')['providers'][0],disabled=True))
        self.assertEqual(member.api('/api/timetable/ai/routes')['models'],[])
        with self.assertRaises(ValueError):member.api('/api/timetable/ai/recognize',{'year':'2026','term':'3','model':identity,'files':files})


class PreviewAndHistoryTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.accounts=AccountStore(self.root/'accounts.sqlite3');self.store=AccountDataStore(self.accounts.database,self.root)
    def tearDown(self):self.temp.cleanup()

    def test_ai_independent_routes_grid_missing_fields_and_formula(self):
        replies=[]
        def complete(config,messages,**kwargs):
            replies.append((config,messages));return json.dumps({'courses':[{'name':'数学','weekday':1,'periods':'1-2','weeks':''}], 'warnings':['周次缺失']})
        service=TimetableAIService(self.store,complete=complete)
        config={'label':'课表专线','protocol':'responses','base_url':'https://example.test/v1','api_key':'private-test-key','models':[{'name':'vision','alias':'vision'}],'default':True}
        identity=service.save(config)['id'];self.assertEqual(service.listing()[0]['id'],identity)
        self.assertNotIn('private-test-key',json.dumps(service.listing(private=True)))
        config2=service.save(dict(config,label='第二条'))['id'];self.assertEqual(sum(m['default'] for m in service.listing()),1)
        service.save(dict(service.listing(private=True)[0],default=True))
        book=Workbook();sheet=book.active;sheet.title='课表';sheet['A1']='星期一';sheet.merge_cells('A2:B2');sheet['A2']='数学';sheet['D3']='=1+1';book.create_sheet('其它')
        stream=BytesIO();book.save(stream);files=[file('网格课表.xlsx',stream.getvalue())]
        self.assertEqual(len(service.inspect({'files':files})['files'][0]['sheets']),2)
        result=service.recognize('owner',{'year':'2026','term':'3','model':identity,'files':[dict(files[0],sheets=['课表'])]})
        self.assertFalse(result['ready']);self.assertTrue(result['issues']);self.assertIn('A2:B2',replies[0][1][0]['content']);self.assertNotIn('=1+1',replies[0][1][0]['content'])
        result['courses'][0]['weeks']='1-16周(单)';validated=normalize_preview(result);self.assertTrue(validated['ready']);self.assertEqual(validated['schedule']['courses'][0]['weeks'],list(range(1,17,2)))
        self.assertFalse(normalize_preview({'year':'2026','term':'3','courses':[]})['ready'])
        overlap=normalize_preview({'year':'2026','term':'3','courses':validated['courses']*2});self.assertTrue(overlap['warnings'])
        csv_file=file('课程.csv','课程名称,星期,节次,周次\n英语,周二,3-4,1-16'.encode())
        parsed=service.recognize('owner',{'year':'2026','term':'3','model':'','files':[csv_file]});self.assertTrue(parsed['ready']);self.assertFalse(parsed['used_ai'])
        with self.assertRaises(ValueError):service.recognize('owner',{'year':'2026','term':'3','model':'missing','files':[dict(files[0],sheets=['课表'])]})
        service.save({'action':'delete','id':config2});self.assertNotIn(b'private-test-key',(self.root/'accounts.sqlite3').read_bytes())

    def test_history_limits_conflict_recovery_and_attachment_metadata(self):
        history=ChatHistory(self.store);user={'role':'user','content':'first message','attachments':[{'name':'a.txt','mime':'text/plain','size':4,'sha256':'a'*64}]}
        value={'request_id':'a'*12,'model':'test:model','version':0}
        conversation,_=history.begin('owner',value,user,'fingerprint','run-id');cid=conversation['id']
        history.checkpoint('owner',value['request_id'],'中断前的内容')
        recovered=ChatHistory(self.store).get('owner',cid);self.assertEqual(recovered['status'],'cancelled');self.assertEqual(recovered['messages'][-1]['content'],'中断前的内容')
        with self.assertRaises(ConversationConflict):history.rename('owner',cid,{'title':'new','version':0})
        with patch('chat_history.MAX_CONVERSATIONS',1):
            with self.assertRaises(HistoryQuota):history.begin('owner',dict(value,request_id='b'*12),user,'other','other')
        with patch('chat_history.MAX_BYTES',1):
            with self.assertRaises(HistoryQuota):history.begin('other',dict(value,request_id='c'*12),user,'other','other')
        self.assertNotIn('中断前的内容'.encode(),(self.root/'accounts.sqlite3').read_bytes())

    def test_global_and_account_concurrency_and_canonical_attachment_context(self):
        chat=ChatService(self.store)
        async def slow(*args,**kwargs):
            yield {'type':'text_delta','text':'partial'}
            await asyncio.sleep(30)
        try:
            config={'label':'test','protocol':'openai','base_url':'https://example.test/v1','api_key':'test','models':[{'name':'x','alias':'x'}]}
            identity=chat.save(config)['id']
            with patch('web_chat.stream_model',slow):
                payload={'model':identity+':x','messages':[{'role':'user','content':'test'}],'request_id':'first-request'}
                one=chat.open_stream('owner',payload)
                self.assertEqual(next(one.events())['event'],'start')
                # Closing an event iterator intentionally cancels its run; keep the second run open.
                one.wait(5)
                two=chat.open_stream('second',dict(payload,request_id='second-request'))
                three=chat.open_stream('third',dict(payload,request_id='third-request'))
                with self.assertRaises(ChatBusy):chat.open_stream('second',dict(payload,request_id='another-request'))
                with self.assertRaises(ChatBusy):chat.open_stream('fourth',dict(payload,request_id='fourth-request'))
                two.cancel();three.cancel();self.assertTrue(two.wait());self.assertTrue(three.wait())
            raw=file('private.txt',b'attachment-text')
            _,metadata=chat._new_message({'messages':[{'role':'user','content':'','attachments':[raw]}]})
            previous={'messages':[dict(metadata,id='saved')]}
            missing=chat._context(previous,[{'role':'user','content':'FORGED'},{'role':'user','content':'next'}],{'role':'user','content':'next'},'openai')
            self.assertNotIn('FORGED',missing[0]['content']);self.assertIn('需要重新上传',missing[0]['content'])
            restored=chat._context(previous,[{'role':'user','content':'FORGED','attachments':[raw]},{'role':'user','content':'next'}],{'role':'user','content':'next'},'openai')
            self.assertIn('attachment-text',restored[0]['content'])
        finally:chat.close()


class ProtocolAdapterTests(unittest.TestCase):
    def test_unicode_multiline_sse_errors_and_idle_timeout(self):
        original=httpx.AsyncClient
        config={'protocol':'openai','base_url':'https://fixture.test/v1','api_key':'test'}
        class Chunked(httpx.AsyncByteStream):
            async def __aiter__(self):
                data=': keepalive\n\ndata: {"choices": [\ndata: {"delta":{"content":"中文"},"finish_reason":null}]}\n\ndata: {"choices":[{"delta":{},"finish_reason":"stop"}]}\n\n'.encode()
                for byte in data:yield bytes([byte])
        async def consume(transport):
            with patch('chat_runtime.httpx.AsyncClient',side_effect=lambda **kwargs:original(transport=transport,**kwargs)):
                return [e async for e in stream_model(config,'x',[{'role':'user','content':'x'}],timeout=.5,idle_timeout=.01)]
        events=asyncio.run(consume(httpx.MockTransport(lambda request:httpx.Response(200,headers={'Content-Type':'text/event-stream'},stream=Chunked()))))
        self.assertEqual(events[0]['text'],'中文')
        for response in [httpx.Response(200,text='{"choices":[]}'),httpx.Response(200,headers={'Content-Type':'text/event-stream'},text='data: {"choices":[{"delta":{"content":"partial"}}]}\n\n')]:
            with self.assertRaises(ValueError):asyncio.run(consume(httpx.MockTransport(lambda request:response)))
        def timeout(request):raise httpx.ReadTimeout('idle')
        with self.assertRaisesRegex(ValueError,'超时'):asyncio.run(consume(httpx.MockTransport(timeout)))


class TrendsTests(unittest.TestCase):
    def test_parsers_and_shared_cache_backoff_stale_expiry(self):
        self.assertEqual(_baidu_rows({'data':{'cards':[{'content':[{'content':[{'word':'百度'}]}]}]}})[0]['word'],'百度')
        self.assertEqual(len(_items('weibo',[{'word':'广告','is_ad':1},{'word':'热搜','num':23}])),1)
        with tempfile.TemporaryDirectory() as root:
            now=[100000.];counts={key:0 for key in SOURCES};gate=threading.Event();peak=[0];active=[0];lock=threading.Lock();fail=[False]
            def fetch(source):
                with lock:counts[source]+=1;active[0]+=1;peak[0]=max(peak[0],active[0])
                gate.wait(2)
                with lock:active[0]-=1
                if fail[0]:raise ValueError('HTTP 403')
                return _items(source,[{'word':'测试热搜','hot_value':1,'hotScore':2,'num':3,'heat_score':4}])
            store=TrendsService(root,fetcher=fetch,clock=lambda:now[0]);self.addCleanup(store.close)
            self.assertTrue(all(s['status']=='loading' for s in store.get()['sources']));store.get(force=True);gate.set()
            for _ in range(100):
                if not store.inflight:break
                time.sleep(.01)
            self.assertEqual(peak[0],2);self.assertTrue(all(v==1 for v in counts.values()))
            self.assertTrue(all(s['status']=='fresh' for s in store.get()['sources']))
            now[0]+=61;fail[0]=True;store.get(force=True)
            for _ in range(100):
                if not store.inflight:break
                time.sleep(.01)
            self.assertTrue(all(s['status']=='stale' and s['items'] for s in store.get()['sources']))
            now[0]+=86401
            self.assertTrue(all(not s['items'] for s in store.get()['sources']))


if __name__=='__main__':unittest.main()
