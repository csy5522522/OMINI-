"""Full browser acceptance in isolated local service; never targets production."""
from io import BytesIO
import json
from pathlib import Path
import time
from PIL import Image
from openpyxl import Workbook
from playwright.sync_api import sync_playwright, expect
import tests.test_web_chat as fixtures


def main():
    dest=Path(__file__).resolve().parents[1]/'artifacts/verification/workspace-v3';dest.mkdir(parents=True,exist_ok=True)
    fixtures.WebChatTests.setUpClass();suite=fixtures.WebChatTests
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch();context=browser.new_context(viewport={'width':1440,'height':1000});page=context.new_page();errors=[]
            page.on('pageerror',lambda error:errors.append(str(error)))
            trends={'poll_after':60,'sources':[{'id':source,'title':name,'url':url,'status':'fresh','updated_at':time.time(),'items':[{'rank':i+1,'title':f'{name} · 今日热点 {i+1}','url':url,'heat':10000-i*203,'tag':''} for i in range(10)]} for source,name,url in [('douyin','抖音热搜','https://www.douyin.com/hot'),('baidu','百度热搜','https://top.baidu.com'),('weibo','微博热搜','https://weibo.com'),('bilibili','哔哩哔哩热搜','https://search.bilibili.com')]]}
            page.route('**/api/trends*',lambda route:route.fulfill(json=trends))
            page.goto(suite.url+'/admin');page.locator('#admin-password').fill('admin123');page.locator('#admin-login button').click();expect(page.locator('#admin-panel')).to_be_visible()
            page.locator('#model-add').click();page.locator('#provider-label').fill('测试助理');page.locator('#provider-base').fill(suite.base);page.locator('#provider-key').fill('fixture-key');page.locator('#provider-models').fill('fixture-model');page.locator('#provider-form [type=submit]').click();expect(page.locator('.model-card')).to_have_count(1)
            page.locator('#timetable-line-add').click();page.locator('#timetable-line-label').fill('课表专用视觉线路');page.locator('#timetable-line-base').fill(suite.base);page.locator('#timetable-line-key').fill('vision-key');page.locator('#timetable-line-model').fill('fixture-model');page.locator('#timetable-line-default').check();page.locator('#timetable-line-form [type=submit]').click();expect(page.locator('#timetable-lines .model-card')).to_have_count(1)
            for theme in ('light','dark'):
                page.evaluate('(theme)=>window.OminiAppearance.setTheme(theme)',theme)
                page.screenshot(path=str(dest/f'admin-{theme}.png'),full_page=True)
            page.goto(suite.url+'/#chat');expect(page.locator('#chat-model option')).to_have_count(1)
            expect(page.locator('html')).to_have_attribute('data-scale','125')
            page.locator('#chat-files').set_input_files({'name':'资料.txt','mimeType':'text/plain','buffer':'历史附件正文'.encode()});expect(page.locator('#chat-attachments .chat-file-chip')).to_have_count(1)
            page.locator('#chat-input').fill('分析附件');page.locator('#chat-send').click();expect(page.locator('.chat-message.assistant .chat-message-body')).to_have_text('你好，测试回复');expect(page.locator('#chat-send')).to_have_attribute('aria-label','发送消息')
            assert page.locator('.chat-message>strong').count()==0
            assert page.locator('.chat-message.user').evaluate("node=>getComputedStyle(node).backgroundColor")=='rgba(0, 0, 0, 0)'
            expect(page.locator('#chat-status')).to_be_hidden()
            for theme in ('light','dark'):
                page.evaluate('(theme)=>window.OminiAppearance.setTheme(theme)',theme);page.screenshot(path=str(dest/f'chat-{theme}.png'),full_page=True,animations='disabled')
            page.reload();expect(page.locator('#chat-model option')).to_have_count(1);page.locator('#chat-history').click();expect(page.locator('.chat-history-row')).to_have_count(1);page.locator('.chat-history-open').click();expect(page.locator('.chat-message.user .chat-file-chip')).to_contain_text('需重新上传')
            page.locator('#chat-input').fill('SLOW');page.locator('#chat-send').click();expect(page.locator('.chat-message.assistant .chat-message-body').last).to_contain_text('部分回复');page.get_by_role('button',name='停止生成',exact=True).click();expect(page.locator('#chat-send')).to_have_attribute('aria-label','发送消息');expect(page.locator('.chat-message.assistant .chat-message-body').last).to_have_text('部分回复')
            page.locator('#settings-open').click()
            for scale in ('100','110','125'):
                page.locator('[data-appearance-scale]').select_option(scale);expect(page.locator('html')).to_have_attribute('data-scale',scale)
            page.locator('[data-appearance-theme]').select_option('light');expect(page.locator('html')).to_have_attribute('data-theme','light');page.locator('[data-close=settings-dialog]').click()
            page.locator('[data-page=timetable]').click();page.locator('#import-open').click();page.locator('#import-tab-ai').click();expect(page.locator('#ai-model option')).to_have_count(1)
            img=BytesIO();Image.new('RGB',(120,100),'white').save(img,'PNG')
            page.locator('#ai-files').set_input_files({'name':'课表.png','mimeType':'image/png','buffer':img.getvalue()});expect(page.locator('#ai-file-list img')).to_be_visible();page.locator('#ai-recognize').click();expect(page.locator('#ai-preview')).to_be_visible();expect(page.locator('#ai-import-confirm')).to_be_disabled()
            page.locator('#ai-preview-rows [data-field=weeks]').fill('1-16周(单)');page.locator('#ai-check').click();expect(page.locator('#ai-import-confirm')).to_be_enabled();page.screenshot(path=str(dest/'ai-preview-light.png'),full_page=True)
            page.locator('#ai-import-confirm').click();expect(page.locator('#import-dialog')).not_to_be_visible();expect(page.locator('#schedule-grid')).to_contain_text('高等数学')
            # Multi-sheet grid Excel reads merged cell coordinates and still requires preview.
            book=Workbook();sheet=book.active;sheet.title='周课表';sheet['A1']='课程';sheet.merge_cells('B2:C2');sheet['B2']='高等数学';book.create_sheet('备用表');buf=BytesIO();book.save(buf)
            page.locator('#import-open').click();page.locator('#import-tab-ai').click();page.locator('#ai-files').set_input_files({'name':'网格.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':buf.getvalue()});expect(page.locator('.ai-sheet-choice')).to_have_count(2);page.locator('#ai-recognize').click();expect(page.locator('#ai-preview')).to_be_visible();expect(page.locator('#ai-import-confirm')).to_be_disabled();page.locator('[data-close=import-dialog]').click()
            page.locator('[data-page=trends]').click();expect(page.locator('.trends-card')).to_have_count(4)
            page.screenshot(path=str(dest/'trends-light.png'),full_page=True)
            page.locator('#theme-toggle').click();expect(page.locator('html')).to_have_attribute('data-theme','dark');page.screenshot(path=str(dest/'trends-dark.png'),full_page=True)
            for width,height in [(1024,768),(390,844)]:
                page.set_viewport_size({'width':width,'height':height})
                for route in ('chat','trends','timetable'):
                    page.locator(f'[data-page={route}]').click();expect(page.locator('#page-'+route)).to_be_visible();assert page.evaluate('document.documentElement.scrollWidth<=innerWidth'),f'overflow {route} at {width}'
                    page.screenshot(path=str(dest/f'{route}-{width}.png'),full_page=True)
            expect(page.locator('html')).to_have_attribute('data-scale','100')
            page.locator('#mobile-settings').click();page.locator('[data-appearance-scale]').select_option('125');expect(page.locator('html')).to_have_attribute('data-scale','125');page.locator('[data-close=settings-dialog]').click();page.screenshot(path=str(dest/'mobile-125.png'),full_page=True)
            assert page.evaluate('document.documentElement.scrollWidth<=innerWidth'),page.evaluate("({width:innerWidth,scroll:document.documentElement.scrollWidth,offenders:[...document.querySelectorAll('body *')].filter(e=>!e.closest('[hidden],dialog:not([open])')&&e.getBoundingClientRect().right>innerWidth+2).slice(0,12).map(e=>({tag:e.tagName,id:e.id,cls:e.className,right:e.getBoundingClientRect().right}))})")
            page.reload();expect(page.locator('html')).to_have_attribute('data-scale','125')
            assert not errors,errors
            browser.close();print('PASS full browser: global themes/scale, streaming/cancel/history, AI preview/grid Excel, four trends/mobile. Screenshots:',dest)
    finally:suite.tearDownClass()


if __name__=='__main__':main()
